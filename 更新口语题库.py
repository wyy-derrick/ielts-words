#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
更新口语题库.py：把 口语题库.xlsx 生成 speaking-questions.js（网页用的题库数据）。

用法（在仓库目录下）：
    /home/wyy/anaconda3/bin/python 更新口语题库.py

说明：
- 只收录「地区」不是「非大陆」的题；新题/保留题/万年老题写入 status，供网页「优先抽新题」用。
- 生成文件勿手改；改题请编辑 口语题库.xlsx 后重跑本脚本。
- 依赖：openpyxl。
"""

import hashlib
import os
import re
import sys
from datetime import datetime

from openpyxl import load_workbook

XLSX_NAME = "口语题库.xlsx"
JS_NAME = "speaking-questions.js"

STATUS_MAP = {"新题": "new", "保留题": "retained", "万年老题": "classic"}


def slug(text):
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s or "topic"


def status_code(value):
    return STATUS_MAP.get(str(value or "").strip(), "retained")


def cell_lines(value):
    return [ln.strip() for ln in str(value or "").split("\n") if ln.strip()]


def read_sheet(wb, prefix):
    for ws in wb.worksheets:
        if ws.title.startswith(prefix):
            rows = ws.iter_rows(values_only=True)
            header = [str(c or "").strip() for c in next(rows, [])]
            return header, list(rows)
    return None, []


def col(header, name):
    return header.index(name) if name in header else None


def pick(row, idx):
    if idx is None or idx >= len(row):
        return ""
    return row[idx]


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(here, XLSX_NAME)
    if not os.path.exists(path):
        sys.exit("没找到 {}，先跑 解析题库PDF.py 生成。".format(XLSX_NAME))
    wb = load_workbook(path, read_only=True)

    # 版本信息在「说明」页（该页没有表头行，逐行找）
    version = ""
    if "说明" in wb.sheetnames:
        for row in wb["说明"].iter_rows(values_only=True):
            if row and str(pick(row, 0)).strip() == "题库版本":
                version = str(pick(row, 1)).strip()

    p1, p2 = [], []
    seen = {}

    h, rows = read_sheet(wb, "P1")
    ti, ri, si, qi = col(h, "话题"), col(h, "地区"), col(h, "状态"), col(h, "小问（一行一个）")
    for row in rows:
        title = str(pick(row, ti) or "").strip()
        if not title or str(pick(row, ri)).strip() == "非大陆":
            continue
        questions = cell_lines(pick(row, qi))
        if not questions:
            print("跳过：P1「{}」没有小问".format(title))
            continue
        item = {"id": "p1-" + slug(title), "title": title,
                "status": status_code(pick(row, si)), "questions": questions}
        if item["id"] in seen:
            print("跳过：P1「{}」编号重复".format(title))
            continue
        seen[item["id"]] = title
        p1.append(item)

    h, rows = read_sheet(wb, "P2")
    ti = col(h, "中文题名")
    ri, si = col(h, "地区"), col(h, "状态")
    ci, pi, p3i, ki = col(h, "英文题卡"), col(h, "提示点（一行一个）"), col(h, "P3追问（一行一个）"), col(h, "剧本")
    for row in rows:
        title = str(pick(row, ti) or "").strip()
        if not title or str(pick(row, ri)).strip() == "非大陆":
            continue
        cue = str(pick(row, ci) or "").strip()
        if not cue:
            print("跳过：P2「{}」没有英文题卡".format(title))
            continue
        item = {
            "id": "p2-" + hashlib.md5(title.encode("utf-8")).hexdigest()[:8],
            "title": title, "status": status_code(pick(row, si)),
            "cue": cue, "points": cell_lines(pick(row, pi)),
            "p3": cell_lines(pick(row, p3i)),
            "script": str(pick(row, ki) or "").strip(),
        }
        if item["id"] in seen:
            print("跳过：P2「{}」编号重复".format(title))
            continue
        seen[item["id"]] = title
        p2.append(item)

    data = {
        "version": version,
        "generatedAt": datetime.now().strftime("%Y-%m-%d"),
        "p1": p1,
        "p2": p2,
    }

    import json
    out = os.path.join(here, JS_NAME)
    with open(out, "w", encoding="utf-8") as f:
        f.write("// 由 更新口语题库.py 从 口语题库.xlsx 生成，勿手改；改题请编辑 xlsx 后重跑。\n")
        f.write("window.__SPEAKING_QUESTIONS__ = ")
        f.write(json.dumps(data, ensure_ascii=False, indent=2))
        f.write(";\n")

    n1_new = sum(1 for t in p1 if t["status"] == "new")
    n2_new = sum(1 for t in p2 if t["status"] == "new")
    print("版本：{}".format(version or "（说明页没写）"))
    print("P1 话题：{} 个进网页（新题 {}）".format(len(p1), n1_new))
    print("P2 题卡：{} 张进网页（新题 {}）".format(len(p2), n2_new))
    print("已生成 {}".format(out))


if __name__ == "__main__":
    main()
